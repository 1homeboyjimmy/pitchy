"""Strict CSV/XLSX reader for accelerator application imports."""

import csv
import io
from itertools import islice
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

MAX_FILE_BYTES = 5_000_000
MAX_UNCOMPRESSED_BYTES = 25_000_000
MAX_ROWS = 2_000
MAX_COLUMNS = 100


def _clean(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()[:10_000]


def parse_form_export(content: bytes, filename: str, selected_sheet: str = "") -> dict:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise HTTPException(422, "Поддерживаются только CSV и XLSX")
    if not content or len(content) > MAX_FILE_BYTES:
        raise HTTPException(422, "Файл пуст или превышает 5 МБ")
    if suffix == ".csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = content.decode("cp1251")
            except UnicodeDecodeError as exc:
                raise HTTPException(422, "Не удалось определить кодировку CSV") from exc
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        matrix = list(csv.reader(io.StringIO(text), dialect))
        sheets = ["CSV"]
        sheet = "CSV"
    else:
        try:
            with ZipFile(io.BytesIO(content)) as archive:
                if "[Content_Types].xml" not in archive.namelist() or sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                    raise HTTPException(422, "Некорректный или слишком большой XLSX")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except (BadZipFile, InvalidFileException, ValueError, OSError) as exc:
            raise HTTPException(422, "Не удалось прочитать XLSX") from exc
        sheets = workbook.sheetnames
        sheet = selected_sheet or sheets[0]
        if sheet not in sheets:
            workbook.close()
            raise HTTPException(422, "Лист XLSX не найден")
        worksheet = workbook[sheet]
        if (worksheet.max_column or 0) > MAX_COLUMNS or (worksheet.max_row or 0) > MAX_ROWS + 1:
            workbook.close()
            raise HTTPException(422, "Слишком много строк или столбцов; максимум 2000 ответов и 100 столбцов")
        matrix = [list(row) for row in islice(worksheet.iter_rows(values_only=True), MAX_ROWS + 2)]
        workbook.close()
    if len(matrix) < 2:
        raise HTTPException(422, "В файле нет строк с ответами")
    if len(matrix) > MAX_ROWS + 1:
        raise HTTPException(422, "Слишком много строк; максимум 2000 ответов")
    headers = [_clean(value) for value in matrix[0]]
    if not headers or len(headers) > MAX_COLUMNS or any(not header for header in headers) or len(headers) != len(set(headers)):
        raise HTTPException(422, "Заголовки столбцов должны быть заполнены и не повторяться")
    numbered_rows = [(number, [_clean(value) for value in row[:len(headers)]] + [""] * max(0, len(headers) - len(row))) for number, row in enumerate(matrix[1:], start=2)]
    numbered_rows = [(number, row) for number, row in numbered_rows if any(row)]
    rows = [row for _, row in numbered_rows]
    if not rows:
        raise HTTPException(422, "В файле нет заполненных ответов")
    return {"headers": headers, "rows": rows, "row_numbers": [number for number, _ in numbered_rows], "sheets": sheets, "sheet": sheet}
